"""Scrape SwimSwam Daily Swim Coach Workouts into one Word document.

URL sources (in order):
  1. urls.txt, if it contains real http(s) lines
  2. Otherwise, auto-generates URLs for workout numbers START_NUM..END_NUM
     and silently skips the ones that don't exist (404).
"""
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

INPUT_FILE = "urls.txt"
OUTPUT_DOCX = "SwimSwam_488_Workouts.docx"
FAILED_FILE = "failed_urls.txt"

START_NUM = 101
END_NUM = 1134
URL_TEMPLATE = "https://swimswam.com/daily-swim-coach-workout-{}/"

MAX_WORKERS = 4
REQUEST_DELAY = 0.5
RETRIES = 3

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}


def load_urls():
    urls = []
    if os.path.exists(INPUT_FILE):
        with open(INPUT_FILE, "r", encoding="utf-8") as f:
            urls = [l.strip() for l in f if l.strip().startswith("http")]
    if urls:
        print(f"Loaded {len(urls)} URLs from {INPUT_FILE}")
        return urls
    print(
        f"No usable URLs in {INPUT_FILE}; generating workouts "
        f"{START_NUM}-{END_NUM} (missing ones will be skipped)."
    )
    return [URL_TEMPLATE.format(n) for n in range(START_NUM, END_NUM + 1)]


def fetch_workout(url):
    match = re.search(r"workout-(\d+)", url)
    num = int(match.group(1)) if match else 0
    last_error = "unknown error"

    for attempt in range(1, RETRIES + 1):
        time.sleep(REQUEST_DELAY)
        try:
            r = requests.get(url, headers=HEADERS, timeout=20)
        except requests.RequestException as e:
            last_error = str(e)
        else:
            if r.status_code == 404:
                return {"num": num, "url": url, "status": "notfound"}
            if r.status_code == 200:
                return parse_page(num, url, r.text)
            last_error = f"HTTP {r.status_code}"
        time.sleep(2 * attempt)  # back off before retrying

    return {"num": num, "url": url, "status": "failed", "error": last_error}


def parse_page(num, url, html):
    soup = BeautifulSoup(html, "html.parser")
    title_tag = soup.find("h1")
    title = title_tag.get_text(strip=True) if title_tag else f"Workout #{num}"

    content = soup.find("div", class_="entry-content") or soup.find(
        "div", class_="post-content"
    )
    if not content:
        return {
            "num": num,
            "url": url,
            "status": "failed",
            "error": "Content div not found",
        }

    for tag in content.find_all(["script", "style", "ins", "iframe", "form"]):
        tag.decompose()

    lines = []
    for el in content.find_all(["p", "ul", "ol", "h2", "h3", "h4"]):
        txt = el.get_text(separator="\n", strip=True)
        if txt:
            lines.append(txt)

    if not lines:
        return {
            "num": num,
            "url": url,
            "status": "failed",
            "error": "Page had no text",
        }

    return {
        "num": num,
        "url": url,
        "status": "ok",
        "title": title,
        "body": "\n\n".join(lines),
    }


def build_docx(good):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.75)
        s.left_margin = s.right_margin = Inches(0.75)

    cover = doc.add_paragraph("SwimSwam Daily Swim Workouts Collection")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover.runs[0]
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102)

    sub = doc.add_paragraph(f"Compiled Collection of {len(good)} Workouts")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(14)
    sub.runs[0].font.italic = True

    for item in good:
        doc.add_page_break()
        h = doc.add_heading(item["title"], level=1)
        h.paragraph_format.space_before = Pt(0)
        h.paragraph_format.space_after = Pt(4)

        p_link = doc.add_paragraph()
        p_link.paragraph_format.space_after = Pt(12)
        r = p_link.add_run(f"Source: {item['url']}")
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = RGBColor(100, 100, 100)

        for para in item["body"].split("\n\n"):
            p = doc.add_paragraph(para)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15

    doc.save(OUTPUT_DOCX)


def main():
    urls = load_urls()
    print(f"Scraping {len(urls)} URLs...")

    results = []
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futures = [ex.submit(fetch_workout, u) for u in urls]
        for i, fut in enumerate(as_completed(futures), start=1):
            res = fut.result()
            results.append(res)
            if res["status"] == "failed":
                print(f"[{i}/{len(urls)}] #{res['num']} FAILED: {res['error']}")
            elif i % 25 == 0 or i == len(urls):
                print(f"[{i}/{len(urls)}] processed")

    results.sort(key=lambda r: r["num"])
    good = [r for r in results if r["status"] == "ok"]
    failed = [r for r in results if r["status"] == "failed"]
    missing = [r for r in results if r["status"] == "notfound"]

    print(
        f"\nSuccess: {len(good)} | Not found (404): {len(missing)} "
        f"| Failed: {len(failed)}"
    )

    if failed:
        with open(FAILED_FILE, "w", encoding="utf-8") as f:
            for r in failed:
                f.write(f"{r['url']}\t{r['error']}\n")
        print(f"Failed URLs written to {FAILED_FILE}")

    if not good:
        print("ERROR: No workouts were scraped. No document created.")
        print("If errors say HTTP 403, the site is blocking this server.")
        print("Run the script on your own computer instead.")
        sys.exit(1)

    build_docx(good)
    print(f"Saved {len(good)} workouts to '{OUTPUT_DOCX}'.")


if __name__ == "__main__":
    main()
